# ทะเบียนบิลที่ส่งบัญชีแล้ว — เดือนไหน · บริษัท · เลขบิล · PO · มูลค่า (user สั่ง 15/10/2569)
#   python -B scripts/acct-sent-register.py   ← เขียนทับ Desktop\ยอดคงคลัง\ยอดคงคลังส่งบัญชี\_ทะเบียนบิลส่งบัญชี.csv
# แหล่ง (อ่านอย่างเดียว): ชีทรายเดือนใน ยอดคงคลัง_รวมปีงบ69.xlsx + Master_ก.ย.69 / Master ใน ยอดคลังยา_69.xlsm
#   แถว B = "ส่งแล้ว" + ชนิดรายการ ซื้อยา* → รวมเป็นบิล (บริษัท+เลขบิล+ชนิด) · มูลค่า = AA (มูลค่าซื้อยาเดือนนั้น)
#   ชีทก่อน พ.ย.68 ไม่มีชนิดรายการ/มูลค่าซื้อ → ยอดรับเข้า × ราคา · ชีทไม่มี PO → หาจากชีท รับยา ด้วยเลขบิล
# ตรวจ: ผลรวมต่อเดือนต้อง "ตรง" กับ AA ทั้งชีท · รันใหม่หลังส่งบัญชี (/verify-bill-v2) · ขึ้นเดือนแล้วเพิ่มชีทใน SOURCES
import openpyxl, warnings, csv, io, os, datetime
from collections import OrderedDict
warnings.filterwarnings('ignore')
H = r'C:\Users\PRH0000484'
FY = H + r'\OneDrive\ยอดคงคลัง_รวมปีงบ69.xlsx'
STOCK = H + r'\OneDrive\ยอดคลังยา_69.xlsm'
RECV = H + r'\OneDrive\รับจากการซื้อ_ยืม_ตุลา2567-2569.xlsm'
OUT = H + r'\Desktop\ยอดคงคลัง\ยอดคงคลังส่งบัญชี\_ทะเบียนบิลส่งบัญชี.csv'
FY_MONTHS = ['ก.ย.68', 'ต.ค.68', 'พ.ย.68', 'ธ.ค.68', 'ม.ค.69', 'ก.พ.69', 'มี.ค.69', 'เม.ย.69', 'พ.ค.69', 'มิ.ย.69', 'ก.ค.69', 'ส.ค.69']
SOURCES = [(FY, m, m) for m in FY_MONTHS] + [(STOCK, 'Master_ก.ย.69', 'ก.ย.69'), (STOCK, 'Master', 'ต.ค.69')]


def bkey(v):
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def col(h, pred):
    for i, v in enumerate(h):
        if v is not None and pred(str(v).strip()):
            return i
    return None


def ymd(v):
    if isinstance(v, datetime.datetime):
        return f'{v.day:02d}/{v.month:02d}/{v.year + 543}'
    return '' if v is None else str(v)


def num(v):
    return float(v) if isinstance(v, (int, float)) else 0.0


RCV = {}
for r in openpyxl.load_workbook(RECV, read_only=True, data_only=True)['รับยา'].iter_rows(min_row=2, values_only=True):
    if r[7] in (None, '', '-'):
        continue
    e = RCV.setdefault(bkey(r[7]), dict(po=set(), dates=set()))
    if r[8] not in (None, '', '-'):
        e['po'].add(bkey(r[8]))
    if isinstance(r[14], datetime.datetime):
        e['dates'].add(r[14])

bills, check, books = [], [], {}
for f, sn, month in SOURCES:
    wb = books.setdefault(f, openpyxl.load_workbook(f, read_only=True, data_only=True))
    if sn not in wb.sheetnames:
        check.append((month, f'ไม่มีชีท {sn}'))
        continue
    rows = list(wb[sn].iter_rows(values_only=True))
    h = rows[0]
    c = dict(
        B=col(h, lambda s: s == 'ส่งบัญชี'), G=col(h, lambda s: s == 'ชนิด'), J=col(h, lambda s: s in ('ราคา/หน่วย', 'ราคา(บาท)/หน่วย')),
        M=col(h, lambda s: s == 'ชนิดรายการ'), O=col(h, lambda s: s == 'วันที่รับเข้า'),
        P=col(h, lambda s: s == 'ปริมาณ (เข้า)' or s.startswith('ยอดรับยาเข้า')),
        Q=col(h, lambda s: s in ('บริษัท', 'บริษัทยา')), R=col(h, lambda s: s == 'เลขที่บิลซื้อ'), F=col(h, lambda s: s.startswith('รหัส')),
        AA=col(h, lambda s: s.startswith('มูลค่าซื้อยา')), PO=col(h, lambda s: 'PO' in s and 'เลขที่' in s))
    old = c['M'] is None  # layout ก่อน พ.ย.68
    miss = [k for k, v in c.items() if v is None and k not in ('PO',) + (('M', 'O', 'AA') if old else ())]
    if miss:
        check.append((month, f'หาคอลัมน์ไม่เจอ {miss}'))
        continue
    grp, total = OrderedDict(), 0.0
    for r in rows[1:]:
        if r[c['F']] in (None, ''):
            continue
        aa = num(r[c['P']]) * num(r[c['J']]) if old else num(r[c['AA']])
        total += aa
        if str(r[c['B']]).strip() != 'ส่งแล้ว' or (not old and not str(r[c['M']]).startswith('ซื้อยา')) or (old and aa <= 0):
            continue
        kind = 'มิใช่ยา' if 'มิใช่ยา' in str(r[c['G']]) else 'ยา'
        g = grp.setdefault((str(r[c['Q']]).strip(), bkey(r[c['R']]), kind), dict(po=set(), dates=set(), n=0, aa=0.0, pj=0.0))
        po = r[c['PO']] if c['PO'] is not None else None
        if po not in (None, '', '-', 0):
            g['po'].add(bkey(po))
        if not old:
            g['dates'].add(ymd(r[c['O']]))
        g['n'] += 1
        g['aa'] += aa
        g['pj'] += num(r[c['P']]) * num(r[c['J']])
    s_aa = sum(g['aa'] for g in grp.values())
    check.append((month, ('[layout เก่า มูลค่า=รับเข้า×ราคา] ' if old else '')
                  + f'{len(grp)} บิล · มูลค่าส่งบัญชี {s_aa:,.2f} · AA ทั้งชีท {total:,.2f} · {"ตรง" if abs(s_aa - total) < 0.01 else "ไม่ตรง"}'))
    for (sup, bill, kind), g in grp.items():
        notes = [] if abs(g['aa'] - g['pj']) < 0.01 else [f"P×J={g['pj']:.2f}"]
        po, dates = sorted(g['po']), sorted(d for d in g['dates'] if d)
        e = RCV.get(bill)
        if not po and e and e['po']:
            po = sorted(e['po']); notes.append('PO จากชีทรับยา')
        if not dates and e and e['dates']:
            dates = [ymd(d) for d in sorted(e['dates'])]; notes.append('วันรับจากชีทรับยา')
        if not po:
            notes.append('ไม่พบ PO')
        bills.append([month, kind, sup, bill, ', '.join(po), ', '.join(dates), g['n'], f"{g['aa']:.2f}",
                      ' · '.join(notes), f'{os.path.basename(f)} ชีท {sn}'])

buf = io.StringIO()
w = csv.writer(buf, lineterminator='\r\n')
w.writerow(['เดือนที่ส่งบัญชี', 'ชนิด', 'บริษัท', 'เลขที่บิล', 'เลขที่ PO', 'วันที่รับเข้า', 'จำนวนรายการ', 'มูลค่า (บาท)', 'หมายเหตุ', 'แหล่งข้อมูล'])
w.writerows(bills)
open(OUT, 'wb').write(('\ufeff' + buf.getvalue()).encode('utf-8'))
for m, t in check:
    print(m, '|', t)
print('บิล', len(bills), '->', OUT)
