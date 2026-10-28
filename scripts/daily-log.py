# บันทึกการทำงานคลังยารายวัน → Excel รายเดือน — user สั่ง 28/10/2569
# เก็บไว้ไล่ดูว่าแต่ละวันทำอะไร + ไฟล์งานที่ทำอยู่ที่ไหน · ไม่พิมพ์/ไม่ส่งทุกวัน รอ user เรียกส่งหัวหน้าทีเดียว
# ใช้: python scripts/daily-log.py [--month 10/2569]   (ไม่ใส่ = เดือนนี้)
# สร้างไฟล์ใหม่ทั้งไฟล์จากทะเบียนทุกครั้ง (รันซ้ำได้) — ห้ามแก้มือในไฟล์ผลลัพธ์ จะหายรอบหน้า
# ดึงจาก: ทะเบียนใบเบิก · ชีท 'รับยา' · ทะเบียนส่งบิลจัดซื้อ · ทะเบียนเอกสารตรวจรับ · ชีท 'รพ.ยืมยา' · ไฟล์ข้อสังเกต
#        · ไฟล์ที่ทำ: ทะเบียนพิมพ์ (PDF ทุกใบ) · ดัชนีบิล (รูปบิล/COA/ตรวจรับ/ยืม-คืน) · รูปใบเบิก
#        + เรื่องที่ไม่มีทะเบียน (ตัดสินใจ/เปลี่ยนวิธีทำงาน/รอใคร) เขียนเองใน _บันทึกเพิ่ม.csv
import csv, datetime, os, sys, warnings
from collections import OrderedDict, defaultdict

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.stdout.reconfigure(encoding='utf-8')
warnings.filterwarnings('ignore')

HOME = r'C:\Users\PRH0000484'
REQ_REG = os.path.join(HOME, r'Desktop\รวมเบิกยา\_ทะเบียนใบเบิก.csv')
HANDOVER = os.path.join(HOME, r'Desktop\ยอดคงคลัง\ส่งบิลให้พี่โก้\_ทะเบียนส่งบิลจัดซื้อ.csv')
INSPECT_ACK = os.path.join(HOME, r'Desktop\ยอดคงคลัง\เอกสารตรวจรับ\_ทะเบียนเอกสารตรวจรับ.csv')
OBS = os.path.join(HOME, r'Desktop\บิลยา_ตรวจรับ\_ข้อสังเกตบิลและใบเบิก.csv')
REQ_PDF = os.path.join(HOME, r'Desktop\รวมเบิกยา\PDF')
PRINT_LOG = os.path.join(REQ_PDF, '_สถานะพิมพ์.csv')
BILL_INDEX = os.path.join(HOME, r'Desktop\บิลยา_ตรวจรับ\_ดัชนีบิล.csv')
RECV_XLSM = os.path.join(HOME, r'OneDrive\รับจากการซื้อ_ยืม_ตุลา2567-2569.xlsm')
OUT_DIR = os.path.join(HOME, r'Desktop\ยอดคงคลัง\บันทึกการทำงาน')
EXTRA = os.path.join(OUT_DIR, '_บันทึกเพิ่ม.csv')
EXTRA_HEAD = ['วันที่', 'หัวข้อ', 'รายละเอียด', 'แจ้งหัวหน้า']
OURS = 'รพ.ประชาธิปัตย์'
TH_MON = ['ม.ค.', 'ก.พ.', 'มี.ค.', 'เม.ย.', 'พ.ค.', 'มิ.ย.', 'ก.ค.', 'ส.ค.', 'ก.ย.', 'ต.ค.', 'พ.ย.', 'ธ.ค.']


def read_csv(path):
    if not os.path.exists(path):
        return []
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def parse_be(s):
    try:
        d, m, y = (int(x) for x in str(s).strip().split('/'))
        return datetime.date(y - 543, m, d)
    except (ValueError, TypeError):
        return None


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    return parse_be(v)


def be(d):
    return f'{d.day:02d}/{d.month:02d}/{d.year + 543}'


def money(x):
    return f'{x:,.2f}'


def num(v):
    try:
        return float(str(v).replace(',', ''))
    except (ValueError, TypeError):
        return 0.0


def count_by(items):
    c = OrderedDict()
    for k in items:
        c[k] = c.get(k, 0) + 1
    return ' · '.join(f'{k} {n}' if n > 1 else k for k, n in c.items())


args = sys.argv[1:]
today = datetime.date.today()
if '--month' in args:
    m, y = args[args.index('--month') + 1].split('/')
    month, year = int(m), int(y) - 543
else:
    month, year = today.month, today.year


def in_month(d):
    return d is not None and d.month == month and d.year == year


# day -> list of (ประเภทงาน, เลขเอกสาร, หน่วยงาน/บริษัท, รายการ, จำนวน, มูลค่า, หมายเหตุ)
detail = defaultdict(list)
summary = defaultdict(list)   # day -> บรรทัดงานที่ทำ
to_head = defaultdict(list)   # day -> เรื่องที่ขอให้หัวหน้าทราบ/ตัดสิน

# ---- ใบเบิก ----
req_by_day = defaultdict(list)
for r in read_csv(REQ_REG):
    d = parse_be(r['วันที่ลง'])
    if not in_month(d) or r.get('สถานะ', 'ใช้') not in ('', 'ใช้'):
        continue
    req_by_day[d].append(r)
    n = int(num(r['จำนวนแถว']))
    detail[d].append(('ลงใบเบิก', r['เลขรับ'], r['หน่วยงาน'], r['ประเภท'] + (f" (ใบลงวันที่ {r['วันที่ใบเบิก']})" if r['วันที่ใบเบิก'] != r['วันที่ลง'] else ''), n, None, r.get('หมายเหตุ', '')))
for d, rs in req_by_day.items():
    rows = sum(int(num(r['จำนวนแถว'])) for r in rs)
    nos = [r['เลขรับ'] for r in rs]
    rng = nos[0] if len(nos) == 1 else f'{nos[0]}–{nos[-1].split("/")[-1]}'
    summary[d].append(f"ลงใบเบิก {len(rs)} ใบ {rows} รายการ ({count_by(r['หน่วยงาน'] for r in rs)}) · {rng}")

# ---- รับยาเข้าคลัง (ชีท รับยา) ----
wb = openpyxl.load_workbook(RECV_XLSM, read_only=True, data_only=True)
bills = OrderedDict()   # (day, สถานะการซื้อ, บิล, บริษัท) -> [รายการ, มูลค่า]
for r in wb['รับยา'].iter_rows(min_row=2, max_col=18, values_only=True):
    d = as_date(r[14])
    if not in_month(d) or not r[3]:
        continue
    kind = str(r[4] or '').strip() or '-'
    key = (d, kind, str(r[7] or '-').strip(), str(r[17] or '-').strip())
    b = bills.setdefault(key, [[], 0.0])
    b[0].append(f'{r[3]} {num(r[9]):g} {r[10] or ""}'.strip())
    b[1] += num(r[13])
recv_by_day = defaultdict(lambda: defaultdict(lambda: [0, 0, 0.0]))
for (d, kind, no, co), (items, val) in bills.items():
    detail[d].append((f'รับยา ({kind})', no, co, ' · '.join(items), len(items), round(val, 2), ''))
    t = recv_by_day[d][kind]
    t[0] += 1; t[1] += len(items); t[2] += val
for d, kinds in recv_by_day.items():
    parts = []
    for kind, (nb, ni, val) in kinds.items():
        if kind == 'การซื้อ':
            parts.append(f'บิลซื้อ {nb} บิล {ni} รายการ {money(val)} บาท')
        else:
            parts.append(f'{kind} {ni} รายการ')
    summary[d].append('รับยาเข้าคลัง: ' + ' · '.join(parts))

# ---- ยืม-คืนยา (ชีท รพ.ยืมยา) ----
for r in wb['รพ.ยืมยา'].iter_rows(min_row=2, max_col=18, values_only=True):
    if not r[6]:
        continue
    borrower, lender = str(r[1] or '').strip(), str(r[3] or '').strip()
    item = f'{r[6]} {num(r[9]):g} {r[10] or ""}'.strip()
    seq = f'ทะเบียนลำดับ {r[0]}'
    d = as_date(r[13])
    if in_month(d):
        what = f'ยืมจาก {lender}' if borrower == OURS else f'ให้ {borrower} ยืม'
        detail[d].append(('ยืมยา', str(r[2] or '-'), lender if borrower == OURS else borrower, item, num(r[9]), None, seq))
        summary[d].append(f'{what}: {item}')
    d = as_date(r[15])
    if in_month(d):
        what = f'คืน {lender}' if borrower == OURS else f'{borrower} คืนยา'
        detail[d].append(('คืนยา (ยืม)', str(r[16] or '-'), lender if borrower == OURS else borrower, item, num(r[9]), None, seq))
        summary[d].append(f'{what}: {item}')
wb.close()

# ---- ส่งบิลให้จัดซื้อ ----
hand = OrderedDict()
for r in read_csv(HANDOVER):
    d = parse_be(r['วันที่ส่ง'])
    if not in_month(d):
        continue
    h = hand.setdefault((d, r['เลขที่ใบ']), [set(), 0.0])
    h[0].add(r['เลขที่บิล']); h[1] += num(r['มูลค่า(บาท)'])
for (d, no), (bs, val) in hand.items():
    detail[d].append(('ส่งบิลให้จัดซื้อ', no, 'งานจัดซื้อ', ', '.join(sorted(bs)), len(bs), round(val, 2), ''))
    summary[d].append(f'ส่งบิลให้จัดซื้อ {no} ({len(bs)} บิล {money(val)} บาท)')

# ---- เอกสารสรุปตรวจรับ ----
ack = OrderedDict()
for r in read_csv(INSPECT_ACK):
    d = parse_be(r['วันที่ออก'])
    if not in_month(d) or r['ไฟล์'].startswith('ยกเลิก') or r.get('หมายเหตุ', '').startswith('ยกเลิก'):
        continue
    a = ack.setdefault((d, r['เลขที่ใบ']), [set(), 0.0, r['วันที่ตรวจรับ']])
    a[0].add(r['เลขที่บิล']); a[1] += num(r['มูลค่า(บาท)'])
for (d, no), (bs, val, insp) in ack.items():
    detail[d].append(('เอกสารสรุปตรวจรับ', no, 'ผู้ตรวจรับ', f'ตรวจรับวันที่ {insp} · ' + ', '.join(sorted(bs)), len(bs), round(val, 2), ''))
    summary[d].append(f'ออกเอกสารสรุปตรวจรับ {no} (ตรวจรับ {insp} · {len(bs)} บิล {money(val)} บาท)')

# ---- ข้อสังเกต (ความผิดปกติที่เจอตอนลงข้อมูล) ----
for r in read_csv(OBS):
    d = parse_be(r['วันที่ตรวจ'])
    if not in_month(d):
        continue
    detail[d].append(('ข้อสังเกต', r['เลขที่'], r['เอกสาร'], f"{r['รายการ']} — {r['ประเภท']}", None, None, r.get('สถานะ', '')))
    to_head[d].append(f"{r['รายการ']} ({r['เอกสาร']} {r['เลขที่']}): {r['ประเภท']}")

# ---- ไฟล์ที่ทำ ----
files = defaultdict(OrderedDict)   # day -> {path: ประเภท}


def pdf_kind(p):
    for key, kind in (('ส่งบิลให้พี่โก้', 'PDF ใบส่งบิลจัดซื้อ'), ('เอกสารตรวจรับ', 'PDF สรุปตรวจรับ'),
                      ('ยืม-คืนยา', 'PDF ใบยืม/ใบคืนยา'), ('เปลี่ยน', 'PDF เปลี่ยน/คืนบริษัท')):
        if key in p:
            return kind
    return 'PDF ใบเบิก'


for r in read_csv(PRINT_LOG):
    d = parse_be(str(r['วันที่ทำ PDF'])[:10])
    p = r['ไฟล์ (ใต้ รวมเบิกยา\\PDF)'].strip()
    if in_month(d) and p:
        p = p if os.path.isabs(p) else os.path.join(REQ_PDF, p)
        files[d].setdefault(p, pdf_kind(p))
for r in read_csv(REQ_REG):
    d = parse_be(r['วันที่ลง'])
    if in_month(d) and r.get('สถานะ', 'ใช้') in ('', 'ใช้'):
        for p in (x.strip() for x in r['รูป'].split(';')):
            if p:
                files[d].setdefault(p, 'รูปใบเบิก')
for r in read_csv(BILL_INDEX):
    d = parse_be(r['วันที่รับ'])
    p = r['ไฟล์รูป'].strip()
    if in_month(d) and p:
        doc = r['เอกสาร'].strip()
        files[d].setdefault(p, doc if doc.startswith('รูป') else f'รูป{doc}')

# ---- บันทึกเพิ่ม (เขียนเอง) ----
os.makedirs(OUT_DIR, exist_ok=True)
if not os.path.exists(EXTRA):
    with open(EXTRA, 'w', encoding='utf-8-sig', newline='') as f:
        csv.writer(f, lineterminator='\r\n').writerow(EXTRA_HEAD)
for r in read_csv(EXTRA):
    d = parse_be(r['วันที่'])
    if not in_month(d):
        continue
    detail[d].append(('งานอื่น', '-', '-', f"{r['หัวข้อ']}: {r['รายละเอียด']}", None, None, ''))
    line = f"{r['หัวข้อ']}: {r['รายละเอียด']}"
    (to_head if r.get('แจ้งหัวหน้า', '').strip() in ('ใช่', '✓') else summary)[d].append(line)

# ---- เขียน Excel ----
days = sorted(set(detail) | set(summary) | set(to_head) | set(files))
out = openpyxl.Workbook()
HEAD_FILL = PatternFill('solid', fgColor='DDEBF7')
WRAP = Alignment(wrap_text=True, vertical='top')


def sheet(ws, header, widths):
    ws.append(header)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i); c.font = Font(bold=True); c.fill = HEAD_FILL; c.alignment = WRAP
    ws.freeze_panes = 'A2'


s1 = out.active
s1.title = 'สรุปหัวหน้า'
sheet(s1, ['วันที่', 'งานที่ทำ', 'เรื่องที่ขอให้หัวหน้าทราบ/ตัดสิน', 'ไฟล์ที่ทำ (ดูชีท ไฟล์ที่ทำ)'], [12, 90, 60, 28])
for d in days:
    s1.append([be(d), '\n'.join('• ' + x for x in summary[d]), '\n'.join('• ' + x for x in to_head[d]),
               '\n'.join(f'{k} {n}' for k, n in OrderedDict((k, list(files[d].values()).count(k)) for k in files[d].values()).items())])
s2 = out.create_sheet('รายละเอียด')
sheet(s2, ['วันที่', 'ประเภทงาน', 'เลขเอกสาร', 'หน่วยงาน/บริษัท', 'รายการ', 'จำนวน', 'มูลค่า (บาท)', 'หมายเหตุ'], [12, 18, 18, 30, 70, 9, 13, 50])
for d in days:
    for row in detail[d]:
        s2.append([be(d), *row])
s3 = out.create_sheet('ไฟล์ที่ทำ')
sheet(s3, ['วันที่', 'ประเภท', 'ชื่อไฟล์ (คลิกเปิด)', 'โฟลเดอร์', 'ไฟล์ยังอยู่'], [12, 24, 60, 70, 11])
missing = 0
for d in days:
    for p, kind in files[d].items():
        here = os.path.exists(p)
        missing += not here
        s3.append([be(d), kind, os.path.basename(p), os.path.dirname(p), '✓' if here else 'ไม่พบ'])
        if here:
            c = s3.cell(row=s3.max_row, column=3)
            c.hyperlink = p
            c.style = 'Hyperlink'
for ws in (s1, s2, s3):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
for c in s2['G'][1:]:
    c.number_format = '#,##0.00'

path = os.path.join(OUT_DIR, f'บันทึกงานคลังยา_{TH_MON[month - 1]}{str(year + 543)[-2:]}.xlsx')
try:
    out.save(path)
except PermissionError:
    sys.exit(f'เขียนไม่ได้ — ปิดไฟล์ {path} ใน Excel ก่อนแล้วรันใหม่')
print(f'{path} · {len(days)} วัน · {sum(len(v) for v in detail.values())} เรื่อง · '
      f'ไฟล์ {sum(len(v) for v in files.values())}' + (f' (ไม่พบ {missing} — ย้าย/ลบไปแล้ว)' if missing else ''))
last = max((d for d in days if d <= today), default=None)
if last:
    print(f'\n{be(last)}')
    for x in summary[last]:
        print('  • ' + x)
    for x in to_head[last]:
        print('  ! ' + x)
