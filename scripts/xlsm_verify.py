# ตรวจไฟล์ .xlsm หลังแก้ผ่าน COM (ขั้น 7–8 ของ /excel-com) — เทียบทีละเซลล์ + codepage/โมดูล VBA
# ใช้: python -B scripts/xlsm_verify.py <ไฟล์เดิม> <ไฟล์ที่แก้> [--plan <plan.json จาก location-moves.mjs>]
#   ไม่มี --plan = แสดงส่วนต่างอย่างเดียว (คนตัดสิน) · มี --plan = เซลล์ที่ต่างต้องอยู่ใน D/E ของแถวในแผนเท่านั้น
#   exit 0 = ผ่าน · 1 = ไม่ผ่าน (ห้ามวางกลับ)
import json, sys, warnings, zipfile

import olefile
import openpyxl
from openpyxl.utils import get_column_letter
from oletools import olevba

sys.stdout.reconfigure(encoding='utf-8')
warnings.filterwarnings('ignore')


def cell_diffs(fa, fb):
    a = openpyxl.load_workbook(fa, read_only=True, data_only=True)
    b = openpyxl.load_workbook(fb, read_only=True, data_only=True)
    if a.sheetnames != b.sheetnames:
        return None
    out = []
    for sn in a.sheetnames:
        ra = list(a[sn].iter_rows(values_only=True))
        rb = list(b[sn].iter_rows(values_only=True))
        for i in range(max(len(ra), len(rb))):
            x = ra[i] if i < len(ra) else ()
            y = rb[i] if i < len(rb) else ()
            for j in range(max(len(x), len(y))):
                u = x[j] if j < len(x) else None
                v = y[j] if j < len(y) else None
                if (u or None) != (v or None) and not (isinstance(u, (int, float)) and isinstance(v, (int, float)) and abs(u - v) < 1e-9):
                    out.append((sn, f'{get_column_letter(j + 1)}{i + 1}', u, v))
    return out


def vba_info(fn):
    raw = zipfile.ZipFile(fn).read('xl/vbaProject.bin')
    d = olevba.decompress_stream(bytearray(olefile.OleFileIO(raw).openstream('VBA/dir').read()))
    i = d.find(b'\x03\x00\x02\x00\x00\x00')
    cp = int.from_bytes(d[i + 6:i + 8], 'little')
    mods = {f'{s}|{p}': c for (_, s, p, c) in olevba.VBA_Parser(fn).extract_macros()}
    return cp, mods


def main():
    args = sys.argv[1:]
    if len(args) < 2:
        print(__doc__ or 'usage: xlsm_verify.py <orig> <work> [--plan plan.json]')
        return 1
    orig, work = args[0], args[1]
    allowed = None
    if '--plan' in args:
        plan = json.load(open(args[args.index('--plan') + 1], encoding='utf-8'))
        allowed = set()
        for m in plan['moves']:
            for r in m['rows']:
                allowed |= {('Master', f'D{r}'), ('Master', f'E{r}')}
        for l in plan['list']:
            allowed |= {('list', f'D{l["row"]}'), ('list', f'E{l["row"]}')}

    ok = True
    diffs = cell_diffs(orig, work)
    if diffs is None:
        print('✗ รายชื่อชีทไม่เท่าเดิม')
        return 1
    print(f'เซลล์ที่ต่าง {len(diffs)}')
    for sn, addr, u, v in diffs[:60]:
        flag = '' if allowed is None or (sn, addr) in allowed else '  ✗ ไม่อยู่ในแผน'
        print(f'  {sn}!{addr}: {str(u)[:40]} → {str(v)[:40]}{flag}')
        if flag:
            ok = False
    if len(diffs) > 60:
        print(f'  … อีก {len(diffs) - 60}')
        if allowed is not None and any((sn, a) not in allowed for sn, a, _, _ in diffs[60:]):
            ok = False

    ca, ma = vba_info(orig)
    cb, mb = vba_info(work)
    same = ma == mb
    print(f'VBA codepage {ca} → {cb} · โมดูล {len(ma)} → {len(mb)} · {"โค้ดเหมือนเดิม" if same else "โค้ดต่าง"}')
    if ca != cb or not same:
        ok = False
    if allowed is None and ok:
        print('VBA ผ่าน · ไม่มี --plan — ส่วนต่างเซลล์ข้างบนคนต้องตัดสินเอง')
    else:
        print('✓ ผ่าน' if ok else '✗ ไม่ผ่าน — ห้ามวางกลับ')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
