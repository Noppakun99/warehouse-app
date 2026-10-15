"""รายงานแจ้งการสั่งซื้อยา — วิเคราะห์ + verify (อ่านอย่างเดียว ไม่เขียนไฟล์ Excel)

อ่านจากสำเนาใน workdir ที่ snapshot.ps1 คัดลอกมา (stock.xlsm / recv.xlsm / report.xlsx)
ห้ามอ่านไฟล์ใน OneDrive ตรงๆ — ไฟล์ที่เปิดใน Excel อยู่ openpyxl เปิดไม่ได้ (Permission denied)

    python pr.py check  --workdir W              ตรวจสูตรเรท 3 เดือน เทียบชีทเบิก
    python pr.py plan   --workdir W [--date YYYY-MM-DD]
    python pr.py verify --workdir W              ตรวจ report_out.xlsx หลัง apply.ps1
    python pr.py line   --workdir W [--date ...] [--note CODE=ข้อความ ...] [--to พี่...]
"""
import argparse, datetime as dt, json, os, re, sys, tempfile, warnings
from collections import defaultdict, Counter

warnings.filterwarnings('ignore')
import openpyxl

ORDER_STATUSES = ('หมดสต็อค', 'สั่งเพิ่ม', 'สั่งเพิ่ม ใกล้หมดอายุ')
NOT_USE_TYPES = ('บันทึกเท่านั้น', 'ยาโครงการ')            # ชนิดรายการที่สูตรเบิกรายเดือนตัดออก
NOT_USE_DEPTS = ('เบิกเพื่อแลกเปลี่ยนคืนยา', 'คืนยา')       # หน่วยงานที่สูตรเบิกรายเดือนตัดออก
NOT_REAL_DISPENSE = ('บันทึกเท่านั้น', 'แก้ไขระบบ')         # ไม่ใช่การเบิกจริง ห้ามใช้เป็น "เบิกล่าสุด"
CHECKED, UNCHECKED = '☑', '☐'
TH_MONTH = ['ม.ค.', 'ก.พ.', 'มี.ค.', 'เม.ย.', 'พ.ค.', 'มิ.ย.', 'ก.ค.', 'ส.ค.', 'ก.ย.', 'ต.ค.', 'พ.ย.', 'ธ.ค.']
# คอลัมน์เบิกรายเดือน E–P = 12 เดือนล่าสุดที่ปิดแล้ว เรียงเก่า→ใหม่ — เดือนของคอลัมน์อยู่แถว 2 (จัดชีทใหม่ 04/10/2569)
MONTH_COLS = ['E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P']

# รายงานซื้อยา: ชื่อคอลัมน์ในโค้ด/plan.json = ตำแหน่งเดิม (Avg/วัน = 'T', หมายเหตุ = 'S')
# 14/10/2569 ย้าย Avg/วัน มาต่อ Avg/เดือน (N) → ตำแหน่งจริง O, ของเดิม O–S เลื่อนเป็น P–T · Sheet1 = ตำแหน่งจริง +1 คอลัมน์
# apply.ps1 มีตารางเดียวกัน ($PHYS) — แก้ตำแหน่งคอลัมน์ต้องแก้ทั้ง 2 ไฟล์
PHYS_MAP = {'O': 'P', 'P': 'Q', 'Q': 'R', 'R': 'S', 'S': 'T', 'T': 'O'}

# ยาซื้อ → ยาสนับสนุน สสจ. ที่เป็นยาเดียวกัน ขนาด/หน่วยเดียวกัน แต่แยกรหัส (หัวหน้านับคงเหลือรวมทั้ง 2 รหัส — user สั่ง 15/10/2569)
# ชีทวิเคราะห์รวมเฉพาะแถวรหัสเดียวกัน ยาสนับสนุนที่ใช้รหัสเดียวกับยาซื้อจึงนับอยู่แล้ว ไม่ต้องใส่ที่นี่
# ค่า = (รหัสยาสนับสนุน, ตัวคูณแปลงหน่วยยาสนับสนุน → หน่วยยาซื้อ) · คู่ที่ user ยืนยันว่าไม่ใช่ยาเดียวกัน/ไม่บวก:
# ยาหอมเทพจิตร ซอง↔ขวด · Clotrimazole 5g↔ลามาโซน 10g (user 15/10/2569: ไม่บวก ใส่หมายเหตุพอ)
SUPPORT_TWIN = {
    '1470011': ('81', 1),     # ขมิ้นชัน 500mg 100เม็ด ↔ 100แคปซูล
    '1520044': ('82', 1),     # ไพลครีม 30g หลอด
    '1520054': ('84', 1),     # ลูกประคบ ลูก
    '1520023': ('64', 0.5),   # มะขามแขก 200เม็ด ↔ 100แคปซูล = ครึ่งกล่อง (user 15/10/2569: แคปซูล = เม็ด ยาเดียวกัน)
}


def twin_q(d, code):
    """(คงเหลือรวม, Avg/วันรวม, รหัสยาสนับสนุน, คงเหลือยาสนับสนุนในหน่วยของยาซื้อ) — ไม่มีคู่ = ค่าของรหัสเดียว"""
    r, ac = d.analysis[code], d.ac
    q, day = num(r[ac['q']]), num(r[ac['day']])
    sc, f = SUPPORT_TWIN.get(code, (None, 1))
    s = d.analysis.get(sc) if sc else None
    if s is None:
        return q, day, None, 0
    sq = num(s[ac['q']]) * f
    return q + sq, day + num(s[ac['day']]) * f, sc, sq


def phys(key):
    return PHYS_MAP.get(key, key)


def out(*a):
    print(*a, flush=True)


def col_idx(letter):
    n = 0
    for ch in letter:
        n = n * 26 + ord(ch) - 64
    return n - 1


def ftext(v):
    return v.text if hasattr(v, 'text') else ('' if v is None else str(v))


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


def be(d):
    return f"{d.day}/{d.month}/{(d.year + 543) % 100:02d}"


def ce(d):
    return f"{d.day}/{d.month}/{d.year}"


def qfmt(v):
    v = float(v or 0)
    return str(int(v)) if v == int(v) else f"{v:g}"


def find_header(rows, needle, limit=12):
    for i, r in enumerate(rows[:limit]):
        if r and needle in [str(x).strip() for x in r if x is not None]:
            return i
    raise SystemExit(f'หาแถวหัวตารางที่มี "{needle}" ไม่เจอ')


def hidx(h, name, prefix=False):
    for i, x in enumerate(h):
        x = '' if x is None else str(x).strip()
        if x == name or (prefix and x.startswith(name)):
            return i
    raise SystemExit(f'หาคอลัมน์ "{name}" ไม่เจอ')


class Data:
    def __init__(self, wd, need_formulas=False):
        self.wd = wd
        p = lambda n: os.path.join(wd, n)
        for n in ('stock.xlsm', 'recv.xlsm', 'report.xlsx'):
            if not os.path.exists(p(n)):
                raise SystemExit(f'ไม่พบ {p(n)} — รัน snapshot.ps1 ก่อน')
        wb = openpyxl.load_workbook(p('stock.xlsm'), read_only=True, data_only=True)
        self._analysis(wb)
        self._master(wb)
        self._dispense(wb)
        if need_formulas:
            self.wbf = openpyxl.load_workbook(p('stock.xlsm'), read_only=True)
        self._receive(p('recv.xlsm'))

    def _analysis(self, wb):
        rows = list(wb['วิเคราะห์สั่งซื้อ'].iter_rows(values_only=True))
        hi = find_header(rows, 'รหัส')
        h = rows[hi]
        self.ah = h
        c = {k: hidx(h, n, pre) for k, n, pre in [
            ('code', 'รหัส', 0), ('type', 'ชนิด', 0), ('name', 'รายการยา', 0), ('unit', 'หน่วย', 0),
            ('total', 'รวมทุกเดือน', 0), ('max', 'Max/เดือน', 1), ('avg', 'Avg/เดือน', 1), ('day', 'Avg/วัน', 1),
            ('q', 'คงเหลือ (หน่วย)', 0), ('ss', 'Safety Stock', 0), ('lt', 'Lead Time', 1), ('rop', 'ROP', 1),
            ('status', 'สถานะ', 0), ('rec', 'จำนวนแนะนำสั่งซื้อ', 1), ('price', 'ราคา/หน่วย', 1),
            ('supplier', 'บริษัทล่าสุด', 0), ('ae', 'ตัดออกจากบัญชี', 1), ('sum3', 'รวมเบิก3เดือนล่าสุด', 0)]}
        self.ac = c
        self.amonth = rows[hi - 1]  # แถว 2 = วันที่ต้นเดือนของคอลัมน์ E–P
        self.arow0 = hi + 2  # เลขแถว Excel ของข้อมูลแถวแรก
        self.analysis = {}
        for r in rows[hi + 1:]:
            if r[c['code']] is None or r[c['type']] is None:
                continue
            self.analysis[str(r[c['code']]).strip()] = r

    def _master(self, wb):
        rows = list(wb['Master'].iter_rows(values_only=True))
        hi = find_header(rows, 'รหัส_txt')
        h = rows[hi]
        ic, iu, ip, iq, ie = (hidx(h, 'รหัส_txt'), hidx(h, 'หน่วย'), hidx(h, 'packsize'),
                              hidx(h, 'คงเหลือหลังจ่าย'), hidx(h, 'Exp'))
        self.pack, self.maxexp = {}, {}
        for r in rows[hi + 1:]:
            if r[ic] is None:
                continue
            k = str(r[ic]).strip()
            self.pack.setdefault(k + str(r[iu]), r[ip])
            if num(r[iq]) > 0 and isinstance(r[ie], dt.datetime):
                self.maxexp[k] = max(self.maxexp.get(k, r[ie]), r[ie])

    def packsize(self, code, unit):
        try:
            return float(self.pack.get(code + str(unit)) or 1) or 1
        except (TypeError, ValueError):
            return 1

    def _dispense(self, wb):
        rows = list(wb['เบิก '].iter_rows(values_only=True))
        hi = find_header(rows, 'วันที่เบิก')
        h = rows[hi]
        c = {k: hidx(h, n) for k, n in [('date', 'วันที่เบิก'), ('code', 'รหัส_txt'), ('kind', 'ชนิดรายการ'),
                                        ('qty', 'ปริมาณ (ออก)'), ('unit', 'หน่วย'), ('dept', 'หน่วยงานที่เบิก'),
                                        ('base', 'qty_base'), ('req', 'จำนวนที่ขอ')]}
        self.dc = c
        self.drow0 = hi + 2
        self.disp = [r for r in rows[hi + 1:] if isinstance(r[c['date']], dt.datetime)]
        self.disp_last_row = hi + 1 + max((i + 1 for i, r in enumerate(rows[hi + 1:]) if r[c['date']] is not None), default=0)

    def _receive(self, path):
        rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True)['รับยา'].iter_rows(values_only=True))
        h = rows[0]
        ic, ist, idt = hidx(h, 'รหัส'), hidx(h, 'สถานะการซื้อ'), hidx(h, 'วันที่รับ')
        self.buy = defaultdict(list)   # วันที่รับ สถานะ "การซื้อ" เท่านั้น — จับคู่ด้วยรหัสอย่างเดียว ไม่สนหน่วย
        for r in rows[1:]:
            if r[ist] == 'การซื้อ' and isinstance(r[idt], dt.datetime):
                self.buy[str(r[ic]).strip()].append(r[idt])
        self.recv_max = max((d for v in self.buy.values() for d in v), default=None)

    def last_buy(self, code):
        v = self.buy.get(code)
        return max(v) if v else None


def load_report(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    rp, s1 = wb['รายงานซื้อยา'], wb['Sheet1']
    if str(rp['E4'].value).strip() != 'รหัสHosxp':
        raise SystemExit(f'หัวคอลัมน์ E4 ของรายงานไม่ใช่ "รหัสHosxp" (เจอ {rp["E4"].value!r}) — โครงสร้างไฟล์เปลี่ยน หยุดก่อน')
    for key, want in (('N', 'Avg/เดือน'), ('T', 'Avg/วัน'), ('O', 'คงเหลือ'), ('S', 'หมายเหตุ')):
        got = str(rp[f'{phys(key)}4'].value or '').strip()
        if not got.startswith(want):
            raise SystemExit(f'หัวคอลัมน์ {phys(key)}4 ไม่ใช่ "{want}…" (เจอ {got!r}) — ตำแหน่งคอลัมน์รายงานเปลี่ยน แก้ PHYS_MAP ใน pr.py + apply.ps1 ก่อน')
    rows = []
    for r in range(5, rp.max_row + 1):
        code = rp.cell(r, 5).value
        if code in (None, ''):
            continue
        rows.append(dict(row=r, code=str(code).strip(), A=rp.cell(r, 1).value, B=rp.cell(r, 2).value,
                         C=rp.cell(r, 3).value, unit=str(rp.cell(r, 8).value).strip(), name=rp.cell(r, 7).value,
                         S=rp[f'{phys("S")}{r}'].value))
    archive = defaultdict(list)
    s1_last = 2
    for r in range(3, s1.max_row + 1):
        code = s1.cell(r, 6).value
        if code in (None, ''):
            continue
        s1_last = r
        if isinstance(s1.cell(r, 3).value, dt.datetime):
            archive[str(code).strip()].append(s1.cell(r, 3).value)
    head = {k: rp[k].value for k in ('A1', f'{phys("O")}4', f'{phys("P")}4', f'{phys("Q")}4')}
    return rows, archive, s1_last, head


# ---------------------------------------------------------------- check
def month_cols(d):
    """คืน [(letter, (year, month))] ของคอลัมน์ E–P จากแถว 2 (ค่าที่ Excel คำนวณไว้ตอนเซฟ)"""
    first = next(d.wbf['วิเคราะห์สั่งซื้อ'].iter_rows(min_row=d.arow0, max_row=d.arow0))
    res = []
    for L in MONTH_COLS:
        v = d.amonth[col_idx(L)]
        if not isinstance(v, dt.datetime):
            raise SystemExit(f'แถว 2 คอลัมน์ {L}: ไม่ใช่วันที่ต้นเดือน ({v!r}) — โครงสร้างชีทวิเคราะห์เปลี่ยน หยุดก่อน')
        res.append((L, (v.year, v.month)))
    return res, first


def formula_cols(f):
    """ตัวอักษรคอลัมน์ที่สูตรอ้าง (เช่น '=MAX(N4,O4,P4)' → ('N','O','P'))"""
    return tuple(re.findall(r'(?<![A-Z$])([A-Z]{1,2})\d+', ftext(f).lstrip('=')))


def rate_window(d, cols):
    """หน้าต่างเรท = คอลัมน์ที่สูตร Avg/เดือน แถวแรกอ้าง → [(letter, (year, month))]"""
    first = next(d.wbf['วิเคราะห์สั่งซื้อ'].iter_rows(min_row=d.arow0, max_row=d.arow0))
    letters = formula_cols(first[d.ac['avg']].value)
    lookup = dict(cols)
    if len(letters) != 3 or any(L not in lookup for L in letters):
        raise SystemExit(f'สูตร Avg/เดือน อ้างคอลัมน์ {letters} — ไม่ใช่ 3 เดือนใน E–P หยุดก่อน')
    return [(L, lookup[L]) for L in letters]


def cmd_check(a):
    d = Data(a.workdir, need_formulas=True)
    today = parse_date(a.date)
    problems = []
    cols, first = month_cols(d)
    ws = d.wbf['วิเคราะห์สั่งซื้อ']
    out('เดือน E–P (แถว 2):', ' '.join(f"{L}={TH_MONTH[m - 1]}{(y + 543) % 100}" for L, (y, m) in cols))
    # 1) สูตรเรทอ้างคอลัมน์ชุดเดียวกันทุกแถว
    refs = Counter()
    for r in ws.iter_rows(min_row=d.arow0):
        if r[0].value is None:
            continue
        for k in ('max', 'avg', 'sum3'):
            refs[(k, formula_cols(r[d.ac[k]].value))] += 1
    out('คอลัมน์ที่สูตร Max/Avg/รวม3เดือน อ้าง:', dict(refs))
    if len({x for (_, x) in refs}) != 1:
        problems.append('สูตร Max/Avg/รวม3เดือน อ้างคอลัมน์ไม่เหมือนกันทุกแถว')
    win = rate_window(d, cols)
    out('หน้าต่างเรท:', ' + '.join(f"{L}={TH_MONTH[m - 1]}{(y + 543) % 100}" for L, (y, m) in win),
        '(3 เดือนล่าสุดที่ปิดแล้ว — ชีทจัดใหม่ 04/10/2569 ไม่รวมเดือนปัจจุบัน)')
    cur = (today.year, today.month)
    y, m = cols[-1][1]
    age = (cur[0] - y) * 12 + cur[1] - m
    if age != 1:
        problems.append(f'คอลัมน์ P = {TH_MONTH[m - 1]}{(y + 543) % 100} แต่ควรเป็นเดือนก่อนหน้า ({age} เดือนจากเดือนนี้) '
                        f'— ไฟล์ยังไม่ได้คำนวณใหม่ตั้งแต่ขึ้นเดือน? เปิดใน Excel ให้คำนวณแล้วเซฟก่อน')
    # 2) สูตรเบิกรายเดือนต้องอ้างทั้งตาราง เบิกยา[...] (เลิกล็อกช่วงแถว 04/10/2569)
    m = re.search(r"\$B\$\d+:\$B\$(\d+)", ftext(first[col_idx('E')].value))
    if m and int(m.group(1)) < d.disp_last_row:
        problems.append(f'สูตรเบิกรายเดือนครอบถึงแถว {m.group(1)} แต่ชีทเบิกมีข้อมูลถึงแถว {d.disp_last_row} — ยอดหลังจากนั้นไม่ถูกนับ')
    out(f'ชีทเบิกมีข้อมูลถึงแถว {d.disp_last_row} | สูตรครอบ {"ถึงแถว " + m.group(1) if m else "ทั้งตาราง เบิกยา[...]"}')
    # 3) ยอดเบิกรายเดือนเทียบชีทเบิก (ตัวกรองเดียวกับสูตร)
    c = d.dc
    agg = defaultdict(float)
    for r in d.disp:
        if r[c['kind']] in NOT_USE_TYPES or r[c['dept']] in NOT_USE_DEPTS:
            continue
        agg[(str(r[c['code']]).strip(), r[c['date']].year, r[c['date']].month)] += num(r[c['base']])
    bad = Counter()
    ex = []
    n = 0
    win_ix = [col_idx(L) for L, _ in win]
    stat_bad = Counter()
    for code, r in d.analysis.items():
        ps = d.packsize(code, r[d.ac['unit']])
        for L, (y, mo) in cols:
            n += 1
            mine = round(agg.get((code, y, mo), 0) / ps, 2)
            got = num(r[col_idx(L)])
            if abs(mine - got) > 0.011:
                bad[L] += 1
                if len(ex) < 8:
                    ex.append((code, r[d.ac['name']], L, got, mine))
        ms = [num(r[i]) for i in win_ix]
        if abs(num(r[d.ac['avg']]) - round(sum(ms) / 3, 2)) > 0.011:
            stat_bad['Avg/เดือน'] += 1
        if abs(num(r[d.ac['max']]) - max(ms)) > 0.011:
            stat_bad['Max/เดือน'] += 1
        if abs(num(r[d.ac['sum3']]) - round(sum(ms), 2)) > 0.011:
            stat_bad['รวม3เดือน'] += 1
    out(f'ยอดเบิกรายเดือน: เทียบ {n} ช่อง ไม่ตรง {sum(bad.values())} {dict(bad)}')
    for e in ex:
        out('   ', e)
    out(f'Max/Avg/รวม3เดือน เทียบหน้าต่างข้างบน: ไม่ตรง {dict(stat_bad) or 0}')
    if bad:
        problems.append('ยอดเบิกรายเดือนในชีทวิเคราะห์ไม่ตรงกับชีทเบิก')
    if stat_bad:
        problems.append('Max/Avg/รวม3เดือน ไม่ตรงกับหน้าต่างที่สูตรควรหยิบ (ไฟล์ยังไม่ได้คำนวณใหม่?)')
    st = Counter(r[d.ac['status']] for r in d.analysis.values())
    out('สถานะ (V):', dict(st))
    out('ข้อมูลรับยาล่าสุด (การซื้อ):', d.recv_max and d.recv_max.date())
    if problems:
        out('\n!! ปัญหาที่ต้องให้ผู้ใช้ตัดสินก่อนทำต่อ:')
        for x in problems:
            out('  -', x)
        sys.exit(2)
    out('\nOK — สูตรเรทและยอดเบิกตรงกับชีทเบิก')


# ---------------------------------------------------------------- plan
def parse_date(s):
    return dt.datetime.strptime(s, '%Y-%m-%d') if s else dt.datetime.combine(dt.date.today(), dt.time())


def row_values(d, code):
    r, c = d.analysis[code], d.ac
    v = lambda k: r[c[k]]
    return dict(I=v('price'), J=v('supplier'), K=v('lt'), L=v('ss'), M=v('max'), N=v('avg'), O=twin_q(d, code)[0],
                P=v('total'), Q=v('sum3'), R=v('rec'), T=v('day'))


def clean(x):
    if isinstance(x, dt.datetime):
        return x.date().isoformat()
    if isinstance(x, float) and x == int(x):
        return int(x)
    return x


def cmd_plan(a):
    d = Data(a.workdir, need_formulas=True)
    today = parse_date(a.date)
    rows, archive, s1_last, head = load_report(os.path.join(a.workdir, 'report.xlsx'))
    warn = []
    dup = [k for k, n in Counter(x['code'] for x in rows).items() if n > 1]
    if dup:
        warn.append(f'รหัสซ้ำในรายงาน: {dup}')
    # 0) ยกเลิกแจ้ง (--cancel CODE=เหตุผล) → ย้ายไป Sheet1 แบบ ☐ + หมายเหตุ ตามธรรมเนียมเดิมของ Sheet1
    cancel = dict(n.split('=', 1) for n in (a.cancel or []))
    unknown = [k for k in cancel if k not in {x['code'] for x in rows}]
    if unknown:
        raise SystemExit(f'--cancel: ไม่พบรหัส {unknown} ในรายงาน')
    dstr = be(today)
    # 1) มาส่งแล้ว = รับสถานะการซื้อ วันที่ >= วันที่แจ้ง (จับคู่รหัสอย่างเดียว)
    moves, keep, c_upd = [], [], []
    for x in rows:
        if x['code'] in cancel:
            note = f"ยกเลิกแจ้ง {dstr}: {cancel[x['code']]}"
            moves.append(dict(row=x['row'], code=x['code'], name=x['name'], kind='cancel', tick=False,
                              S=f"{x['S']} / {note}" if x['S'] else note, notified=str(x['B'])[:10], delivered=None))
            r = d.analysis.get(x['code'])
            if r is not None and r[d.ac['status']] in ORDER_STATUSES:
                warn.append(f"{x['code']} {x['name']}: ยกเลิกแจ้งแล้ว แต่ชีทวิเคราะห์ (ที่เซฟไว้) ยังเป็น '{r[d.ac['status']]}' "
                            f"— ต้องตั้ง AE เป็น '📋 สั่งเมื่อขอ' แล้วเซฟ ไม่งั้นรอบหน้าจะกลับมาเป็นรายการใหม่")
            continue
        b = x['B']
        lb = d.last_buy(x['code'])
        newc = lb.date().isoformat() if lb else '-'
        oldc = x['C'].date().isoformat() if isinstance(x['C'], dt.datetime) else str(x['C'])
        if newc != oldc:
            c_upd.append(dict(code=x['code'], value=newc, old=oldc))
        if not isinstance(b, dt.datetime):
            warn.append(f"แถว {x['row']} {x['code']}: วันที่แจ้ง (B) ไม่ใช่วันที่ ({b!r}) — ไม่ย้าย")
            keep.append(x)
            continue
        after = [t for t in d.buy.get(x['code'], []) if t >= b]
        # แถว "แจ้งซ้ำ: ของมาแล้ว d/m/yy" ถูกสร้างเพราะการรับวันนั้นเอง — การรับวันเดียวกับวันแจ้งไม่ใช่ของรอบใหม่
        # (เดิมรอบที่ 2 ในวันเดียวกันย้ายแถวแจ้งซ้ำไป Sheet1 แล้วเพิ่มกลับเป็นรายการใหม่วนซ้ำ — 14/10/2569)
        if f"แจ้งซ้ำ: ของมาแล้ว {be(b)}" in str(x['S'] or ''):
            after = [t for t in after if t.date() > b.date()]
        if after:
            moves.append(dict(row=x['row'], code=x['code'], name=x['name'], kind='delivered', tick=x['A'] != CHECKED, S=None,
                              notified=b.date().isoformat(), delivered=min(after).date().isoformat()))
        else:
            if x['A'] == CHECKED:
                warn.append(f"แถว {x['row']} {x['code']} {x['name']}: ติ๊ก ☑ แล้ว แต่ไม่พบการรับ 'การซื้อ' หลังวันแจ้ง — ไม่ย้าย ให้ผู้ใช้ตรวจ")
            keep.append(x)
    # 2) อัพเดตตัวเลขของแถวที่ยังรอของ (จับคู่รหัส+หน่วย)
    refresh = []
    for x in keep:
        r = d.analysis.get(x['code'])
        if r is None:
            warn.append(f"{x['code']} {x['name']}: ไม่มีในชีทวิเคราะห์ — ไม่อัพเดตตัวเลข")
            continue
        if str(r[d.ac['unit']]).strip() != x['unit']:
            warn.append(f"{x['code']} {x['name']}: หน่วยในรายงาน '{x['unit']}' ≠ ชีทวิเคราะห์ '{r[d.ac['unit']]}' — ไม่อัพเดตตัวเลข")
            continue
        refresh.append(dict(code=x['code'], values={k: clean(v) for k, v in row_values(d, x['code']).items()}))
        _, _, sc, sq = twin_q(d, x['code'])
        if sc:
            warn.append(f"{x['code']} {x['name']}: คงเหลือในรายงาน = ยาซื้อ + ยาสนับสนุนรหัส {sc} ({qfmt(sq)}) — หมายเหตุแถวนี้ต้องบอกเอง (สคริปต์ไม่แตะหมายเหตุแถวเดิม)")
    # 3) รายการใหม่ — ห้ามซ้ำกับรายการที่ยังรอของ (แบบ ก)
    pending = {x['code'] for x in keep}
    moved = {m['code']: m for m in moves}
    order = {'หมดสต็อค': 0, 'สั่งเพิ่ม ใกล้หมดอายุ': 1, 'สั่งเพิ่ม': 2}
    new = []
    for code, r in d.analysis.items():
        stv = r[d.ac['status']]
        if stv not in ORDER_STATUSES or code in pending or code in cancel:
            continue
        note = []
        if stv == 'สั่งเพิ่ม ใกล้หมดอายุ' and code in d.maxexp:
            note.append(f"สั่งเพิ่ม ใกล้หมดอายุ Exp.{d.maxexp[code]:%d/%m/%Y}")
        prev = [dt.datetime.fromisoformat(moved[code]['notified'])] if code in moved else archive.get(code, [])
        lb = d.last_buy(code)
        if prev and lb and lb >= max(prev):
            note.append(f"แจ้งซ้ำ: ของมาแล้ว {be(lb)} แต่ยังต่ำกว่า ROP")
        _, _, sc, sq = twin_q(d, code)
        if sc:
            note.append(f"คงเหลือรวมยาสนับสนุน สสจ. (รหัส {sc}) {qfmt(sq)} {r[d.ac['unit']]}")
        new.append(dict(code=code, status=stv, F=r[d.ac['type']], G=r[d.ac['name']], H=r[d.ac['unit']],
                        C=lb.date().isoformat() if lb else '-', S=' / '.join(note),
                        **{k: clean(v) for k, v in row_values(d, code).items()}))
    new.sort(key=lambda x: (order[x['status']], str(x['G']).lower()))
    # 4) หัวรายงาน
    cols, _ = month_cols(d)
    win = [m for _, m in rate_window(d, cols)]
    allm = sorted(m for _, m in cols)
    mm = lambda t: f"{TH_MONTH[t[1] - 1]}{(t[0] + 543) % 100}"
    headers = {
        'A1': [r'UPDATE: วันที่ [\d/]+', f'UPDATE: วันที่ {dstr}'],
        f'{phys("O")}4': [r'คงเหลือ\s+[\d/]+', f'คงเหลือ  {dstr}'],
        f'{phys("P")}4': [r'\(.*\)', f'({mm(allm[0])} - {mm(allm[-1])})'],
        f'{phys("Q")}4': [r'\(.*\)', f'({TH_MONTH[win[0][1] - 1]} - {TH_MONTH[win[-1][1] - 1]} {(win[-1][0] + 543) % 100})'],
    }
    plan = dict(date=today.date().isoformat(), report_rows=len(rows), sheet1_last=s1_last,
                moves=moves, c_updates=c_upd, refresh=refresh, new=new, headers=headers, warnings=warn)
    with open(os.path.join(a.workdir, 'plan.json'), 'w', encoding='utf-8') as f:
        json.dump(plan, f, ensure_ascii=False, indent=1, default=str)
    # preview
    out(f"วันที่แจ้งรอบนี้: {dstr} | รายงานเดิม {len(rows)} แถว")
    dl = [m for m in moves if m['kind'] == 'delivered']
    out(f"\n[1] มาส่งแล้ว → ย้ายไป Sheet1 (☑): {len(dl)} รายการ (ยังไม่ติ๊ก {sum(m['tick'] for m in dl)})")
    for m in dl:
        out(f"   {m['code']:>9} {str(m['name'])[:34]:34} แจ้ง {m['notified']} รับ {m['delivered']}{'  ← ยังไม่ติ๊ก' if m['tick'] else ''}")
    cn = [m for m in moves if m['kind'] == 'cancel']
    if cn:
        out(f"[1b] ยกเลิกแจ้ง → ย้ายไป Sheet1 (☐ + หมายเหตุ): {len(cn)} รายการ")
        for m in cn:
            out(f"   {m['code']:>9} {str(m['name'])[:34]:34} แจ้ง {m['notified']} | S: {m['S']}")
    out(f"\n[2] อัพเดตคอลัมน์ C (วันที่ยาส่งล่าสุด): {len(c_upd)} แถว")
    out(f"[3] อัพเดตตัวเลข I–T ของแถวที่ยังรอของ: {len(refresh)}/{len(keep)} แถว")
    out(f"\n[4] รายการใหม่: {len(new)} {dict(Counter(x['status'] for x in new))} "
        f"(แจ้งซ้ำเพราะของไม่พอ {sum('แจ้งซ้ำ' in x['S'] for x in new)})")
    for x in new:
        out(f"   {x['code']:>9} | {str(x['G'])[:32]:32} | {x['H']} | {x['status']} | คงเหลือ {x['O']} | แนะนำ {x['R']} | {x['S']}")
    out('\n[5] หัวรายงาน:', {k: v[1] for k, v in headers.items()})
    if warn:
        out('\n!! คำเตือน (ต้องแจ้งผู้ใช้):')
        for w in warn:
            out('  -', w)
    out(f"\nบันทึกแผนที่ {os.path.join(a.workdir, 'plan.json')}")


# ---------------------------------------------------------------- verify
def cmd_verify(a):
    wd = a.workdir
    plan = json.load(open(os.path.join(wd, 'plan.json'), encoding='utf-8'))
    src = openpyxl.load_workbook(os.path.join(wd, 'report.xlsx'))
    dst = openpyxl.load_workbook(os.path.join(wd, 'report_out.xlsx'))
    dstv = openpyxl.load_workbook(os.path.join(wd, 'report_out.xlsx'), data_only=True)
    sp, dp, dpv = src['รายงานซื้อยา'], dst['รายงานซื้อยา'], dstv['รายงานซื้อยา']
    errs = []
    moved = [m['code'] for m in plan['moves']]
    srows = {str(sp.cell(r, 5).value).strip(): r for r in range(5, sp.max_row + 1) if sp.cell(r, 5).value not in (None, '')}
    drows = [r for r in range(5, dp.max_row + 1) if dp.cell(r, 5).value not in (None, '')]
    dcodes = [str(dp.cell(r, 5).value).strip() for r in drows]
    exp_codes = [c for c in srows if c not in moved] + [x['code'] for x in plan['new']]
    if Counter(dcodes) != Counter(exp_codes):
        errs.append(f'ชุดรหัสในรายงานไม่ตรงแผน: ขาด {set(exp_codes) - set(dcodes)} เกิน {set(dcodes) - set(exp_codes)}')
    if len(dcodes) != len(set(dcodes)):
        errs.append('มีรหัสซ้ำในรายงาน')
    tref = [t.ref for t in dp.tables.values()]
    last = drows[-1] if drows else 4
    if tref and not tref[0].endswith(str(last)):
        errs.append(f'ตาราง {tref} ไม่ครอบถึงแถวสุดท้าย {last}')
    cmap = {x['code']: x['value'] for x in plan['c_updates']}
    rmap = {x['code']: x['values'] for x in plan['refresh']}
    nmap = {x['code']: x for x in plan['new']}
    d0 = dt.datetime.fromisoformat(plan['date'])

    def same(v, e):
        if isinstance(e, (int, float)) and not isinstance(e, bool):
            return isinstance(v, (int, float)) and abs(v - e) < 1e-9
        if e in (None, ''):
            return v in (None, '')
        if isinstance(v, dt.datetime):
            return v.date().isoformat() == e
        return str(v).strip() == str(e).strip()

    for r in drows:
        code = str(dp.cell(r, 5).value).strip()
        g = lambda L: dp[f'{phys(L)}{r}'].value   # L = ชื่อคอลัมน์ในแผน → ตำแหน่งจริง
        if not str(dp[f'D{r}'].value or '').startswith('='):
            errs.append(f'D{r} ไม่ใช่สูตร')
        if code in nmap:
            x = nmap[code]
            if g('A') != UNCHECKED or not (isinstance(g('B'), dt.datetime) and g('B') == d0):
                errs.append(f'แถวใหม่ {r} {code}: A/B ไม่ถูก')
            for L in 'CFGHIJKLMNOPQRST':
                if not same(g(L), x[L]):
                    errs.append(f'แถวใหม่ {r} {code} {L}: {g(L)!r} ≠ {x[L]!r}')
        else:
            o = srows[code]
            rowless = lambda v: (re.sub(r'([A-Z])\d+', r'\1#', v) if isinstance(v, str) and v.startswith('=')
                                 else (None if v == '' else v))
            for L in 'ABDEFGHS':   # สูตร D เลื่อนเลขแถวตามการลบแถว — เทียบแบบตัดเลขแถวออก
                P = phys(L)
                if rowless(dp[f'{P}{r}'].value) != rowless(sp[f'{P}{o}'].value):
                    errs.append(f'แถวเดิม {code} คอลัมน์ {L} ถูกแก้โดยไม่อยู่ในแผน')
            ec = cmap.get(code)
            if ec is not None and not same(g('C'), ec):
                errs.append(f'{code} C: {g("C")!r} ≠ {ec!r}')
            if ec is None and dp[f'C{r}'].value != sp[f'C{o}'].value:
                errs.append(f'{code} C ถูกแก้โดยไม่อยู่ในแผน')
            for L, e in rmap.get(code, {}).items():
                if not same(g(L), e):
                    errs.append(f'{code} {L}: {g(L)!r} ≠ {e!r}')
    # Sheet1
    s1 = dst['Sheet1']
    start = plan['sheet1_last'] + 1
    got = [str(s1.cell(r, 6).value).strip() for r in range(start, start + len(moved))]
    if got != moved:
        errs.append(f'Sheet1 แถว {start}+ ไม่ตรงกับรายการที่ย้าย')
    for i, m in enumerate(plan['moves']):
        r = start + i
        if m['kind'] == 'delivered' and s1.cell(r, 2).value != CHECKED:
            errs.append(f'Sheet1 แถว {r} {m["code"]}: มาส่งแล้วแต่ไม่ได้ติ๊ก ☑')
        if m['kind'] == 'cancel':
            if s1.cell(r, 2).value == CHECKED:
                errs.append(f'Sheet1 แถว {r} {m["code"]}: ยกเลิกแจ้งต้องเป็น ☐ ไม่ใช่ ☑')
            s_col = col_idx(phys('S')) + 2   # Sheet1 = ตำแหน่งในรายงาน +1 คอลัมน์ (col_idx เริ่ม 0)
            if str(s1.cell(r, s_col).value or '').strip() != m['S'].strip():
                errs.append(f'Sheet1 แถว {r} {m["code"]}: หมายเหตุไม่ตรง ({s1.cell(r, s_col).value!r})')
    if s1.cell(start + len(moved), 6).value not in (None, ''):
        errs.append('Sheet1: มีแถวเกินจากที่ย้าย (ซ้ำ?)')
    for k, (_, rep) in plan['headers'].items():
        if rep not in str(dp[k].value):
            errs.append(f'หัว {k} ไม่มี "{rep}"')
    out(f'รายงาน: {len(drows)} แถว (คาด {len(exp_codes)}) | Sheet1 เพิ่ม {len(moved)} แถวที่ {start}–{start + len(moved) - 1} | ตาราง {tref}')
    if dpv['D5'].value is None and drows:
        out('หมายเหตุ: ค่าแคชของสูตร D ว่าง (ไฟล์ยังไม่ได้คำนวณ) — ปกติถ้า apply.ps1 เซฟผ่าน Excel จะมีค่า')
    if errs:
        out(f'\n!! ไม่ผ่าน {len(errs)} จุด:')
        for e in errs[:40]:
            out('  -', e)
        sys.exit(2)
    out('OK — ผ่านทุกข้อ')


# ---------------------------------------------------------------- line
def cmd_line(a):
    d = Data(a.workdir)
    today = parse_date(a.date)
    rp_path = os.path.join(a.workdir, 'report_out.xlsx')
    if not os.path.exists(rp_path):
        rp_path = os.path.join(a.workdir, 'report.xlsx')
    rows, archive, _, _ = load_report(rp_path)
    notify = {}
    for x in rows:
        if isinstance(x['B'], dt.datetime):
            notify[x['code']] = max(notify.get(x['code'], x['B']), x['B'])
    notes = dict(n.split('=', 1) for n in (a.note or []))
    c = d.dc
    last, unmet = {}, {}
    # ขอเบิกแล้วจ่ายไม่ครบ — รวมต่อคำขอ (วันที่+หน่วยงาน+รหัส+หน่วย) เพราะคำขอเดียวแตกหลายแถว lot ได้
    # จ่าย 0 = "ไม่มีจ่าย" · จ่ายบางส่วน = "จ่ายได้ N ไม่พอ" (เดิมนับเฉพาะจ่าย 0 — แถวที่ได้บางส่วนหลุด 14/10/2569)
    asked = defaultdict(lambda: [0.0, 0.0])
    for r in d.disp:
        code = str(r[c['code']]).strip()
        if r[c['kind']] in NOT_REAL_DISPENSE:
            continue
        if num(r[c['qty']]) > 0:
            key = (r[c['date']], r[c['dept']], r[c['qty']], r[c['unit']], r[c['req']])
            if code not in last or r[c['date']] >= last[code][0]:
                last[code] = key
        k = (r[c['date']], str(r[c['dept']]).strip(), code, str(r[c['unit']]).strip())
        asked[k][0] += num(r[c['qty']])
        asked[k][1] = max(asked[k][1], num(r[c['req']]))
    for (date, dept, code, unit), (got, req) in asked.items():
        if req > 0 and got < req and (code not in unmet or date >= unmet[code][0]):
            unmet[code] = (date, dept, got, unit, req)
    ac = d.ac
    sel = []
    twin_skip = []
    for code, r in d.analysis.items():
        stv = r[ac['status']]
        q, day, sc, sq = twin_q(d, code)
        if sc and stv in ORDER_STATUSES and q > 0 and q >= day * 7:
            twin_skip.append(f"{code} {str(r[ac['name']]).strip()} (รวมยาสนับสนุนรหัส {sc} = {qfmt(q)})")
            continue
        if stv == 'หมดสต็อค' or (stv in ORDER_STATUSES and day > 0 and q < day * 7):
            sel.append((0 if stv == 'หมดสต็อค' else 1, code, r))
    sel.sort(key=lambda t: t[0])
    # รูปแบบที่ user แก้แล้วส่งจริง 14/10/2569 — ไม่มี "เรียน…"/"ขอบคุณ…" · เหตุผลลงท้าย "ครับ" ทุกข้อ
    lines = [f"ฝาก{a.to}ติดตามยา ตามรายการด้านล่างนี้ครับ", ""]
    audit, items = [], []
    for grp, code, r in sel:
        unit = r[ac['unit']]
        sup = str(r[ac['supplier']] or '-').strip()
        sup = sup[len('บริษัท '):] if sup.startswith('บริษัท ') else sup
        n = notify.get(code)
        prev = max(archive[code]) if archive.get(code) else None
        lb = d.last_buy(code)
        tq, tday, sc, sq = twin_q(d, code)
        s = (f"{str(r[ac['name']]).strip()} คงเหลือ {qfmt(tq)} x {unit} ที่คลัง "
             + (f"(รวมยาสนับสนุน สสจ. {qfmt(sq)}) " if sc else "") +
             f"แจ้งสั่งวันที่ {ce(n) if n else (ce(prev) + ' (รอบก่อน)' if prev else '-')} บริษัท {sup} ")
        L = last.get(code)
        s += (f"เบิกล่าสุดวันที่ {ce(L[0])} {str(L[1]).strip()} จำนวน {qfmt(L[2])} x {L[3]} " if L else "ยังไม่มีประวัติเบิก ")
        s += f"วันที่มาส่งล่าสุด {ce(lb) if lb else '-'}"
        # เหตุผล — จากข้อมูลจริงเท่านั้น
        why = []
        q, day = tq, tday
        if q > 0:
            why.append(f"เหลือ {qfmt(q)} {unit} ใช้สัปดาห์ละประมาณ {qfmt(round(day * 7))} {unit} ไม่พอจ่ายสัปดาห์นี้")
        if n and n.date() == today.date():
            got_prev = prev and lb and lb >= prev
            why.append(f"ของรอบที่แจ้ง {be(prev)} มาส่ง {be(lb)} แต่ไม่พอ จึงแจ้งซ้ำวันนี้" if got_prev
                       else "คลังหมดแล้ว เพิ่งแจ้งสั่งวันนี้ รบกวนช่วยเร่งให้ด้วยครับ")
        elif n:
            got = lb and lb >= n
            why.append(f"ของรอบที่แจ้งมาส่ง {be(lb)} แต่ไม่พอ" if got
                       else f"แจ้งไปแล้ว {(today - n).days} วัน ยังไม่มีของเข้า")
        elif prev:
            why.append(f"ของรอบที่แจ้ง {be(prev)} มาส่ง {be(lb)} แต่ไม่พอ ยังไม่มีแถวแจ้งรอบใหม่ในรายงาน" if lb and lb >= prev
                       else "ยังไม่มีแถวแจ้งสั่งในรายงาน")
        else:
            why.append("ยังไม่มีแถวแจ้งสั่งในรายงาน")
        U = unmet.get(code)
        # ด่วน = หน่วยงานขอเบิกภายใน 14 วันแล้วยังไม่ได้ของ/ได้ไม่ครบ (user สั่ง 14/10/2569 ให้ยกขึ้นบนแยกหัวข้อ)
        urgent = bool(U and (not L or U[0] >= L[0]) and (today - U[0]).days <= 14)
        if U and (not L or U[0] >= L[0]):
            why.append(f"วันที่ {ce(U[0])} {U[1]} ขอเบิก {qfmt(U[4])} x {U[3]} "
                       + ("แต่ไม่มีจ่าย" if U[2] <= 0 else f"จ่ายได้ {qfmt(U[2])} ไม่พอจ่าย"))
        if code in notes:
            why.insert(0, notes[code])
        s += f"\n   เหตุผล: {' / '.join(why)} ครับ"
        items.append((0 if urgent else 1, grp, s))
        audit.append((code, r[ac['status']], 'แจ้ง' if n else ('Sheet1' if prev else 'ไม่มี')))
    items.sort(key=lambda t: (t[0], t[1]))
    heads = {0: 'ด่วน — หน่วยงานขอเบิกแล้วยังไม่ได้ของ / ได้ไม่ครบ', 1: 'ติดตามทั่วไป'}
    i = 0
    for sec in (0, 1):
        part = [t for t in items if t[0] == sec]
        if part:
            lines += [f"【{heads[sec]}】", ""]
        for _, _, s in part:
            i += 1
            lines += [f"{i}. {s}", ""]
    out('\n'.join(lines).rstrip())
    out('\n---- ตรวจสอบ (ไม่ต้องส่งไลน์) ----')
    out(f'{len(sel)} รายการ (ด่วน {sum(1 for t in items if t[0] == 0)}): หมดสต็อค {sum(1 for t in sel if t[0] == 0)} · ไม่พอจ่ายภายใน 7 วัน {sum(1 for t in sel if t[0] == 1)}')
    if twin_skip:
        out('ไม่อยู่ในไลน์เพราะยาสนับสนุนรหัสคู่ยังพอ (SUPPORT_TWIN):', '; '.join(twin_skip))
    miss = [c for c, _, src in audit if src != 'แจ้ง']
    if miss:
        out('!! ไม่มีแถวแจ้งสั่งในรายงานปัจจุบัน (ควรเพิ่มก่อนส่ง):', miss)
    # ขอเบิกไม่ครบ 14 วันล่าสุด แต่ไม่ติดไลน์ (คลังมีของ — มักรอตรวจรับ หรือจำกัดจ่ายตั้งใจ) ให้ user ตัดสินเอง
    picked = {code for _, code, _ in sel}
    for code, (date, dept, got, unit, req) in ((k, v) for k, v in unmet.items() if k not in picked and (today - v[0]).days <= 14):
        r = d.analysis.get(code)
        nm = str(r[ac['name']]).strip() if r is not None else code
        st = f"{r[ac['status']]} คงเหลือ {qfmt(r[ac['q']])}" if r is not None else '-'
        out(f'?? ขอเบิกไม่ครบแต่ไม่อยู่ในไลน์: {ce(date)} {dept} {nm} ขอ {qfmt(req)} x {unit} ได้ {qfmt(got)} | {st}')
    unk = [k for k in notes if k not in {c for c, _, _ in audit}]
    if unk:
        out('!! --note ระบุรหัสที่ไม่อยู่ในรายการ:', unk)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('cmd', choices=['check', 'plan', 'verify', 'line'])
    ap.add_argument('--workdir', default=os.path.join(tempfile.gettempdir(), 'purchase-request-report'))
    ap.add_argument('--date', help='วันที่แจ้ง YYYY-MM-DD (ค่าเริ่มต้น = วันนี้)')
    ap.add_argument('--note', action='append', help='เหตุผลเพิ่มจากผู้ใช้ รูปแบบ CODE=ข้อความ (ใช้ซ้ำได้)')
    ap.add_argument('--cancel', action='append', help='plan: ยกเลิกแจ้ง รูปแบบ CODE=เหตุผล → ย้ายไป Sheet1 แบบ ☐ + หมายเหตุ')
    ap.add_argument('--to', default='พี่', help='ชื่อผู้รับในคำขึ้นต้น เช่น "พี่หัวหน้า"')
    a = ap.parse_args()
    if hasattr(sys.stdout, 'reconfigure'):
        sys.stdout.reconfigure(encoding='utf-8')
    {'check': cmd_check, 'plan': cmd_plan, 'verify': cmd_verify, 'line': cmd_line}[a.cmd](a)


if __name__ == '__main__':
    main()
