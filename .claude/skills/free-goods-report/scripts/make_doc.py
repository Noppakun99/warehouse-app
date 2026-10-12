# สร้างแบบรายงานการรับส่วนแถม (.docx) จากไฟล์ต้นแบบ — เขียนทับข้อความใน run เดิม เพื่อคง font/w:cs ของ Word
# ใช้: python make_doc.py jobs.json [--template <docx>] [--outdir <dir>]
# jobs.json:
# { "date": "2026-09-30",
#   "jobs": [ { "out": "สสจ.Paracetamol 500mg_แบบรายงานการรับส่วนแถม_30.9.2026.docx",
#               "source": "สำนักงานสาธารณสุขจังหวัดปทุมธานี",
#               "items": [ { "text": "Acetaminophen 500 mg (Paracetamol) บรรจุ 1 กล่อง/1000 เม็ด จำนวน 1,500 กล่อง",
#                            "lot": "38596608", "exp": "2027-04-23" } ] } ] }
import sys, json, argparse, datetime, os
import docx

sys.stdout.reconfigure(encoding='utf-8')
D = r'C:\Users\PRH0000484\Desktop\ยอดคงคลัง\ใช้\ใบของแถม'
ap = argparse.ArgumentParser()
ap.add_argument('jobs')
ap.add_argument('--template', default=os.path.join(D, 'Unison_แบบรายงานการรับส่วนแถม_27.8.2026.docx'))
ap.add_argument('--outdir', default=D)
a = ap.parse_args()
cfg = json.load(open(a.jobs, encoding='utf-8'))

MONTHS = ['มกราคม', 'กุมภาพันธ์', 'มีนาคม', 'เมษายน', 'พฤษภาคม', 'มิถุนายน',
          'กรกฎาคม', 'สิงหาคม', 'กันยายน', 'ตุลาคม', 'พฤศจิกายน', 'ธันวาคม']
THDIG = str.maketrans('0123456789', '๐๑๒๓๔๕๖๗๘๙')
day = datetime.date.fromisoformat(cfg['date'])
DM = f'{day.day} {MONTHS[day.month - 1]}'          # "30 กันยายน"
YEAR = str(day.year + 543)

def exp_be(iso):                                   # Exp แสดงเป็น พ.ศ. dd/mm/yyyy · ไม่มี exp (ชุดตรวจ/ถุงของ) = "-"
    if not iso or iso == '-':
        return '-'
    e = datetime.date.fromisoformat(iso)
    return f'{e.day:02d}/{e.month:02d}/{e.year + 543}'

def setr(p, i, expect, new):
    r = p.runs[i]
    assert r.text == expect, f'ต้นแบบไม่ตรงที่คาด: run {i} = {r.text!r} (คาด {expect!r}) — ต้นแบบถูกแก้? ดู SKILL.md'
    r.text = new

for j in cfg['jobs']:
    items = j['items']
    d = docx.Document(a.template); P = d.paragraphs
    # หน้า 1 — ใบตรวจรับพัสดุ
    p = P[4]; setr(p, 4, '27 สิงห', DM); setr(p, 5, 'า', ''); setr(p, 6, 'คม', ''); setr(p, 8, '2569', YEAR)
    p = P[5]; setr(p, 6, '1', str(len(items))); setr(p, 11, 'บริษัท ', j['source']); setr(p, 12, 'ยูนีซัน จำกัด', '')
    p = P[6]; setr(p, 1, 'Sodium biphosphate ', ' และ '.join(it['text'] for it in items))
    for r in p.runs[2:]: r.text = ''
    p = P[8]
    setr(p, 3, '27 สิงหาคม 2569............................', f'{DM} {YEAR}............................')
    setr(p, 9, '27 สิงหาคม 2569........................................................', f'{DM} {YEAR}........................................................')
    # หน้า 2 — แบบรายงานการรับส่วนแถม
    p = P[33]; setr(p, 3, '27 สิงห', DM); setr(p, 4, 'า', ''); setr(p, 5, 'คม', ''); setr(p, 6, ' 2569', ' ' + YEAR)
    p = P[35]; setr(p, 2, '27 สิงหาคม 2569.............', f'{DM} {YEAR}.............')
    p = P[36]; setr(p, 2, 'บริษัท ยูนีซัน จำกัด', j['source'])
    lines = [f"{str(n).translate(THDIG)}. {it['text']} Lot: {it['lot']} Exp: {exp_be(it['exp'])}"
             for n, it in enumerate(items, 1)]
    setr(p, 7, '๑.', lines[0].split(' ', 1)[0])          # "๑."
    setr(p, 9, 'Sodium biphosphate 19 g/ 118 ml +Sodium phosphate 7g/ 118 ml ',
         '\n'.join([lines[0].split(' ', 1)[1]] + lines[1:]))   # \n = ขึ้นบรรทัดใหม่ใน run เดียวกัน
    for r in p.runs[10:]: r.text = ''
    body = '\n'.join(x.text for x in d.paragraphs)
    for bad in ('ยูนีซัน', 'UNIMA', 'LE2624', '27 สิงห'):
        assert bad not in body, f'ข้อความต้นแบบหลงเหลือ: {bad}'
    out = os.path.join(a.outdir, j['out'])
    d.save(out); print('saved', out)
