# หาแถว บริจาค/สนับสนุน ในชีท 'รับยา' ของวันที่กำหนด แล้วเทียบกับ Master — อ่านอย่างเดียว ไม่เขียนไฟล์
# ใช้: python find_rows.py 2026-09-30 [--recv <path>] [--master <path>]
import sys, argparse, datetime, warnings
import openpyxl

warnings.filterwarnings('ignore')
sys.stdout.reconfigure(encoding='utf-8')

ap = argparse.ArgumentParser()
ap.add_argument('date')
ap.add_argument('--recv', default=r'C:\Users\PRH0000484\OneDrive\รับจากการซื้อ_ยืม_ตุลา2567-2569.xlsm')
ap.add_argument('--master', default=r'C:\Users\PRH0000484\OneDrive\ยอดคลังยา_69.xlsm')
a = ap.parse_args()
day = datetime.date.fromisoformat(a.date)
KINDS = ('การบริจาค', 'การสนับสนุน')

def as_date(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    return None

# รับยา: B=รหัส C=รูปแบบ D=รายการยา E=สถานะการซื้อ F=Lot G=Exp H=เลขที่บิล/หนังสือ J=จำนวน K=หน่วย L=ราคา/หน่วย M=ราคารวม O=วันที่รับ P=สถานะตรวจรับ R=บริษัท
wb = openpyxl.load_workbook(a.recv, read_only=True, data_only=True)
hits = []
for i, r in enumerate(wb['รับยา'].iter_rows(min_row=2, values_only=True), start=2):
    if as_date(r[14]) == day and r[4] in KINDS:
        hits.append(dict(row=i, code=str(r[1]).strip(), form=r[2], name=r[3], kind=r[4], lot=str(r[5]).strip(),
                         exp=as_date(r[6]), doc=r[7], qty=r[9], unit=r[10], price=r[11], total=r[12],
                         status=r[15], source=r[17]))
others = [i for i, r in enumerate(wb['รับยา'].iter_rows(min_row=2, values_only=True), start=2)
          if as_date(r[14]) == day and r[4] not in KINDS]

# Master: F=รหัส H=รายการยา I=หน่วย K=Lot L=Exp M=ชนิดรายการ O=วันที่รับเข้า P=ปริมาณเข้า Q=บริษัท R=เลขที่บิล
mb = openpyxl.load_workbook(a.master, read_only=True, data_only=True)
master = []
for i, r in enumerate(mb['Master'].iter_rows(min_row=2, values_only=True), start=2):
    master.append(dict(row=i, code=str(r[5]).strip(), name=r[7], unit=r[8], lot=str(r[10]).strip(),
                       exp=as_date(r[11]), kind=r[12], recv=as_date(r[14]), qty=r[15], source=r[16], doc=r[17]))

print(f'วันที่รับ {day} — บริจาค/สนับสนุน {len(hits)} แถว (แถวประเภทอื่นวันเดียวกัน: {others or "-"})\n')
for h in hits:
    print(f"[รับยา!{h['row']}] {h['kind']} | {h['name']} | รหัส {h['code']} | {h['qty']} × {h['unit']} | "
          f"Lot {h['lot']} Exp {h['exp']} | หนังสือ {h['doc']} | จาก {h['source']} | {h['status']} | "
          f"ราคา/หน่วย {h['price']} รวม {h['total']}")
    flags = []
    m = [x for x in master if x['lot'] == h['lot'] and x['recv'] == day and x['name'] == h['name']] \
        or [x for x in master if x['lot'] == h['lot'] and x['recv'] == day]
    if not m:
        flags.append('ไม่พบแถวคู่ใน Master (lot+วันที่รับ)')
    for x in m:
        print(f"    Master!{x['row']}: รหัส {x['code']} | {x['name']} | {x['unit']} | {x['kind']} | {x['qty']}")
        if x['code'] != h['code']:
            flags.append(f"รหัสไม่ตรง: รับยา={h['code']} Master={x['code']} (มักเป็นรหัสชั่วคราว ให้แก้ฝั่งรับยาตาม Master)")
        if x['name'] != h['name']:
            flags.append('ชื่อยาไม่ตรงกับ Master')
    if h['lot'] in ('-', '', 'None'):
        flags.append('ไม่มี lot')
    if not h['doc'] or 'รอ' in str(h['doc']):
        flags.append(f"เลขที่หนังสือ = {h['doc']} (ยังไม่มีเอกสาร — ถามก่อนทำใบ)")
    if h['exp'] and (h['exp'] - day).days < 180:
        flags.append(f"หมดอายุอีก {(h['exp'] - day).days} วัน")
    if isinstance(h['total'], (int, float)) and h['total'] >= 100000:
        flags.append(f"มูลค่ารวมสูง {h['total']:,.2f} บาท — ยืนยันจำนวนกับหนังสือ")
    for f in flags:
        print('    ⚠', f)
    print()
