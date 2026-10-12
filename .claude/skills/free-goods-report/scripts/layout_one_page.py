# จัดหน้าใบของแถมแบบ "แต่ละส่วนอยู่หน้าเดียว" (ผู้ใช้กำหนด 09/10/2569 — ทำทุกใบ ตามใบ 7.10.2026 ที่พิมพ์จริง)
# ใช้กับไฟล์ที่ make_doc.py เพิ่งสร้าง (ย่อหน้ายังเรียงตามต้นแบบ Unison 43 ย่อหน้า) — แก้ไฟล์ผลลัพธ์ ห้ามแก้ต้นแบบ
# ใช้: python layout_one_page.py "<ไฟล์ .docx ในโฟลเดอร์ใบของแถม>"
#   1. หน้า 1 ย่อหน้า "คือ …" (P6) + หน้า 2 รายการตั้งแต่ "๑." (P36 run 7+) → 14pt (sz/szCs 28) ระยะบรรทัด 1.0
#   2. ตัด "\n" นำหน้า "กลุ่มงานเภสัชกรรม…ได้ตรวจรับ" (P37 run 0)
#   3. หัวเรื่องหน้า 2 (P29) → pageBreakBefore
#   4. ลบย่อหน้าว่าง 12 ย่อหน้า: P2 · P7 · P10 · P21–28 (8 บรรทัดที่เคยดันไปหน้าใหม่) · P32 — ต้องว่างและไม่ใช่ anchor กล่องลายเซ็น
import sys
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')
path = sys.argv[1]
d = docx.Document(path)
P = d.paragraphs
assert len(P) == 43, f'คาด 43 ย่อหน้า (ผลจาก make_doc.py) ได้ {len(P)} — จัดหน้าไปแล้ว?'
assert P[6].text.startswith('คือ '), P[6].text[:20]
assert P[29].text.startswith('แบบรายงานการรับส่วนแถม'), P[29].text[:20]
assert P[36].runs[7].text.endswith('.'), P[36].runs[7].text
assert P[37].runs[0].text == '\n' and P[37].text.lstrip('\n').startswith('กลุ่มงานเภสัชกรรม'), repr(P[37].text[:20])


def single_spacing(p):
    pPr = p._p.get_or_add_pPr()
    sp = pPr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing'); pPr.append(sp)
    sp.set(qn('w:after'), '0'); sp.set(qn('w:line'), '240'); sp.set(qn('w:lineRule'), 'auto')


def size14(run):
    rPr = run._r.get_or_add_rPr()
    for tag in ('w:sz', 'w:szCs'):
        el = rPr.find(qn(tag))
        if el is None:
            el = OxmlElement(tag); rPr.append(el)
        el.set(qn('w:val'), '28')


single_spacing(P[6])
for r in P[6].runs:
    if r.text:
        size14(r)
single_spacing(P[36])
for r in P[36].runs[7:]:
    if r.text:
        size14(r)
P[37].runs[0].text = ''
P[29].paragraph_format.page_break_before = True

DROP = [2, 7, 10] + list(range(21, 29)) + [32]
for i in sorted(DROP, reverse=True):
    p = P[i]
    assert p.text.strip() == '' and '<w:drawing' not in p._p.xml, f'P{i} ไม่ว่าง/เป็น anchor: {p.text[:20]!r}'
    p._p.getparent().remove(p._p)
d.save(path)
print('จัดหน้าแล้ว: เหลือ', len(d.paragraphs), 'ย่อหน้า (ใบ 7.10 = 31)')
